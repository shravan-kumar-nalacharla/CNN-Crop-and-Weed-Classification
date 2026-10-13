"""Build the beginner guide from measured project outputs; never use smoke metrics."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def P(text):return ('p',text)
def H(text):return ('h',text)
def T(rows,widths=None):return ('table',rows,widths)
def C(text,language='python'):return ('code',text,language)
def I(name):return ('image',name)
def L(title,url):return ('link',title,url)

def pages_for(r,history,splits):
    pct=lambda x:f'{100*x:.2f}%'
    measured=bool(r)
    summary=(f"The completed five-epoch run achieved {pct(r['test_accuracy'])} test accuracy "
             f"on {r['evaluated_test_images']} held-out images. Its test loss was {r['test_loss']:.4f} "
             f"and ROC AUC was {r['auc']:.4f}." if measured else
             'Internal layout draft. Final test results will be inserted after full training completes.')
    pages=[]
    def add(title,*blocks):pages.append((title,list(blocks)))
    add('Understand Your VGG16 Crop and Weed Project',
        P('A beginner study guide for Nalacherla Shravan Kumar, 23EG110C43, Roll Number 43.'),
        P('This guide assumes that you are starting Deep Learning from the beginning. It explains what the project does, how the model learns, what the files contain, how to read the results, and how to explain the work in class.'),
        P(summary),
        H('How to study this guide'),
        P('First read pages 2–9 to understand the task and the model. Next read pages 10–13 with the results slides open. Then use pages 14–17 to run the project and practise the presentation. Finish with the viva questions and the glossary.'),
        H('The project in one sentence'),
        P('I adapted an ImageNet-pretrained VGG16 convolutional network to classify a seedling photograph as Crop or Weed, then measured its performance on images kept separate from training.'),
        H('What you should be able to explain'),
        P('You should know why the dataset is split, why VGG16 is pretrained, what changes during the two training phases, and what the test metrics mean. You do not need to memorise every line of code before you understand those ideas.'),
        P('Study project completed locally on Windows using Python 3.11 and TensorFlow. Dataset: V2 Plant Seedlings. Execution date: 13 October 2026.'))
    add('1 What Was Built and Why',
        H('The practical problem'),
        P('A crop is a plant grown intentionally, such as maize or wheat. A weed is a plant that is unwanted in a particular growing area. Young plants can look similar because many have small green leaves. This project studies whether a computer can separate the two image categories.'),
        P('The project uses a fixed dataset mapping. Maize, Common Wheat and Sugar Beet are labelled Crop. The other nine dataset species are labelled Weed. This mapping defines the experiment; it is not a universal rule that every plant species is always a weed in every setting.'),
        H('Input and output'),
        P('The input is one RGB seedling image. RGB means that each pixel has red, green and blue values. The output is one number between zero and one: the model’s Weed probability score. A score below 0.5 gives Crop. A score of 0.5 or above gives Weed.'),
        H('What changed in the repository'),
        P('The earlier group project was converted into an individual Roll 43 project using VGG16 transfer learning. I prepared a separate Python environment, downloaded and verified the dataset, created disjoint splits, implemented the model, tested both training phases, ran full-dataset training, and generated results, a report and a stronger presentation.'),
        P('The original course reference documents were retained under references/. The notebook and command-line trainer both call scripts/core.py, so they use the same data processing, architecture and evaluation logic.'),
        P('A repository is a project folder with version history. Git records file changes and GitHub hosts the repository online. The source, documents and measured result artifacts can be reviewed there. Large raw images, the Python environment and the trained model remain local.'),
        H('What this experiment establishes'),
        P('It measures classification performance on the specified held-out image split. It does not establish that an agricultural robot can safely remove plants in a real field. Field backgrounds, lighting and repeated plant identities need additional evaluation.'))
    add('2 Artificial Intelligence Machine Learning and Deep Learning',
        H('The three terms'),
        P('Artificial Intelligence is the broad idea of making computers perform tasks that require useful decision-making or pattern recognition. Machine Learning is one way to do that: the computer learns patterns from examples rather than relying only on manually written rules. Deep Learning is Machine Learning built around neural networks with many processing layers.'),
        H('A simple comparison'),
        P('A manually written image rule might say that a long narrow leaf means Crop. That rule would fail for some grasses and for unusual camera views. A learning model instead receives many labelled examples and adjusts numerical weights to make better predictions.'),
        H('What a neural network contains'),
        P('A neuron is a small numerical operation. It combines input values using weights and usually passes the result through an activation function. A layer contains many such operations. Several layers can transform raw pixels into features and then into a decision. The biological term is an analogy; the model is a mathematical program.'),
        H('Supervised learning'),
        P('This is supervised learning because training images have known labels. Crop is encoded as 0 and Weed as 1. The model predicts a score, compares it with the known label, calculates an error called loss, and updates its weights. The labels provide the supervision.'),
        H('Training and prediction'),
        P('Training changes model weights using labelled examples. Prediction, also called inference, uses already learned weights to classify an image. Predicting one image does not train the model again. The saved .keras file stores the learned model for later inference.'),
        P('A useful analogy is a student learning from practice questions and then answering a new question. The training examples are the practice material. The test set is the final exam that must remain separate.'))
    add('3 How a CNN Reads an Image',
        H('Pixels become numbers'),
        P('An image is a grid of numbers. This project resizes every image to 224 by 224 pixels and uses three colour channels. One image therefore has shape 224 × 224 × 3. A tensor is a multidimensional array of numbers; a batch of four images has shape 4 × 224 × 224 × 3. Resizing creates a consistent shape without adding new information.'),
        H('Convolution and filters'),
        P('A convolutional layer moves small learned filters across the image. A filter is also called a convolution kernel. It combines nearby values and creates a feature map. Early layers can respond to edges or colour changes, and deeper layers combine patterns. These are examples of possible features, not a claim that every filter has a simple human-readable meaning.'),
        H('Why the same filter moves across the image'),
        P('A useful leaf pattern can appear in different locations. Applying shared filter weights across locations lets a CNN detect a pattern without storing a different rule for every pixel position. CNN stands for Convolutional Neural Network.'),
        H('ReLU and pooling'),
        P('ReLU is an activation that replaces negative values with zero while keeping positive values. For example, it changes −2 to 0 and keeps 3 as 3. It introduces nonlinearity. Pooling reduces spatial size. VGG16 uses max pooling: a region containing 2, 4, 1 and 3 would produce 4. These are teaching examples, not model measurements.'),
        H('From feature maps to a decision'),
        P('The convolution base extracts features. A classification head combines those features for the project’s two labels. The head is the part adapted to Crop and Weed. This separation is important because the base starts with features learned from ImageNet.'))
    add('4 VGG16 and Transfer Learning',
        H('What VGG16 is'),
        P('VGG16 is a CNN introduced by Simonyan and Zisserman. Its original thirteen convolutional layers and three fully connected layers give sixteen layers with learned weights. Here, include_top=False removes the ImageNet classifier and keeps the convolution base.'),
        H('What ImageNet contributes'),
        P('ImageNet is a large labelled image dataset. Its learned VGG16 weights provide an existing visual feature extractor. The new head still needs to learn the Crop and Weed decision; pretrained weights alone do not establish this project’s performance.'),
        H('Transfer learning'),
        P('Transfer learning reuses knowledge from another task. Here, a new binary head learns from plant images using pretrained VGG16 features. This avoids learning every visual feature from random weights on a much smaller dataset.'),
        H('Frozen weights'),
        P('Freezing a layer means that training does not update its weights. The frozen layer still processes images and produces features. During phase one, all VGG16 layers are frozen and the new head learns the binary classification task.'),
        H('Fine tuning'),
        P('Fine-tuning allows selected pretrained weights to change slightly for the new task. During phase two, only VGG16 block5 is enabled, while blocks1–4 stay frozen. The classification head remains trainable. A much smaller learning rate protects the useful pretrained features from large changes.'),
        P('VGG16 names the pretrained backbone. The project’s new binary head replaces the original ImageNet head.'),
        H('The five convolution blocks'),
        T([['Block','Conv layers','Filters','Shape after pooling'],['1','2','64','112 × 112 × 64'],['2','2','128','56 × 56 × 128'],['3','3','256','28 × 28 × 256'],['4','3','512','14 × 14 × 512'],['5','3','512','7 × 7 × 512']], [.65,1.1,.85,3.5]),
        P('VGG16 uses 3 × 3 convolution filters. Each block ends in 2 × 2 max pooling. The table shows how spatial size decreases while the number of feature channels increases.'))
    add('5 Understand the Dataset',
        P('The verified V2 Plant Seedlings dataset contains 5,539 RGB images from twelve species. Crop contains 973 images and Weed contains 4,566 images. The Weed class is much larger, so accuracy alone can hide poor recognition of crops.'),
        T([['Species','Class','Images'],['Maize','Crop','257'],['Common Wheat','Crop','253'],['Sugar Beet','Crop','463'],['Black-grass','Weed','309'],['Charlock','Weed','452'],['Cleavers','Weed','335'],['Common Chickweed','Weed','713'],['Fat Hen','Weed','538'],['Loose Silky-bent','Weed','762'],['Scentless Mayweed','Weed','607'],['Shepherd’s Purse','Weed','274'],['Small-flowered Cranesbill','Weed','576']], [3.9,1.15,1.05]),
        H('What was verified'),
        P('The verifier decoded every image and checked the species counts. The downloaded archive contained two species-folder trees. The downloader selected one complete tree so the dataset was not counted twice. It also normalised an incorrectly decoded apostrophe in a folder name.'),
        H('Source and limitation'),
        P('The dataset comes from Kaggle’s V2 copy of the Aarhus Plant Seedlings dataset and uses CC BY-SA 4.0. Images have controlled backgrounds, and multiple growth-stage images can depict the same plant. These facts limit how confidently the result can be transferred to real fields.'))
    add('6 Splitting and Preprocessing',
        H('Training validation and test sets'),
        P('Training images update the weights. Validation images measure progress and select the best checkpoint. Test images measure the selected model’s final performance. Mixing these roles would give an unfair estimate of how well the model handles unseen data.'),
        T([['Partition','Crop','Weed','Total'],['Training','681','3196','3877'],['Validation','146','685','831'],['Test','146','685','831']], [2.8,1.1,1.1,1.1]),
        P('The split is approximately 70% training, 15% validation and 15% test. Seed 42 fixes the random split. Stratification keeps similar Crop and Weed proportions in all three partitions. The code checks that image paths do not overlap between partitions.'),
        H('Data augmentation'),
        P('Augmentation creates small variations during training: horizontal flips, small rotations, zoom and translation. The label remains the same. The split happens before augmentation, and these random transforms are disabled for validation and test images.'),
        H('VGG16 preprocessing'),
        P('Pixel values remain in the original 0–255 range before VGG16 preprocessing. The preprocessing reverses RGB channels to BGR and subtracts ImageNet channel means. Dividing by 255 first would produce the wrong scale for this VGG16 convention. The saved model includes preprocessing so later predictions use the same conversion.'),
        H('One important limit'),
        P('Disjoint image paths do not guarantee disjoint plant identities. Two different images can show the same plant at different growth stages. A stronger future experiment would split using verified plant identities rather than only image paths.'))
    add('7 Your Model Layer by Layer',
        P('The backbone extracts visual features and the new head learns the binary decision. The model has 14,780,481 parameters in total. During the frozen-head phase, 65,793 parameters are trainable and the pretrained convolution parameters remain fixed. A Dense layer is fully connected: each output combines all values from the previous feature vector.'),
        T([['Component','What it does'],['Input','Receives a 224 × 224 × 3 RGB image.'],['Training augmentation','Creates small random training variations. It is inactive during evaluation.'],['VGG16 preprocessing','Converts channel order and subtracts ImageNet means inside the model.'],['VGG16 convolution base','Produces feature maps of shape 7 × 7 × 512 using its thirteen convolution layers.'],['Global average pooling','Averages each feature map, giving a vector of 512 values.'],['Dense 128 with ReLU','Learns combinations of the 512 features for the new task.'],['Dropout 0.5','Randomly zeros about half the head activations during training to reduce reliance on particular units.'],['Dense 1 with sigmoid','Produces a Weed score between zero and one.']], [2.15,3.95]),
        H('Why global average pooling'),
        P('Averaging each map creates a compact feature vector. Flattening every position would instead produce 25,088 values before the dense head. The compact pooling head uses fewer new parameters and is practical for a small transfer-learning dataset.'),
        H('What dropout does'),
        P('Dropout affects feature activations during training. It does not discard half the training images or remove half the stored weights. During normal prediction, dropout is disabled and the head uses its learned features.'),
        H('Interpret the output'),
        P('A Weed score of 0.8 would be classified as Weed because it exceeds 0.5. That is a teaching example, not a measured prediction. A sigmoid score should not be treated as a guaranteed or calibrated confidence level.'))
    add('8 How the Model Learns',
        H('One training step'),
        P('A batch of images enters the model. The model computes predictions in a forward pass. Binary cross-entropy compares the scores with the known labels. Backpropagation calculates how trainable weights contributed to the loss. The Adam optimizer then adjusts those weights. This repeats for all training batches.'),
        P('A gradient describes how loss changes when a weight changes. Backpropagation efficiently computes these gradients through the layers. Adam uses recent gradient information to adapt updates for individual weights, with the learning rate setting the overall update scale.'),
        H('Epoch batch and learning rate'),
        P('An epoch is one pass through the training set. A batch is a small group processed in one step. With 3,877 training images and batch size four, one epoch has 970 training steps; the last batch is smaller. Learning rate controls the size of weight updates. It is a configured setting, not an accuracy value.'),
        T([['Setting','This measured run'],['Head phase','3 epochs, all VGG16 weights frozen'],['Fine-tuning phase','2 epochs, block5 and head trainable'],['Total measured budget','5 epochs on the full training partition'],['Batch size and input','4 images, each 224 × 224 RGB'],['Head learning rate','0.001'],['Fine-tuning learning rate','0.00001, one hundred times smaller'],['Optimizer and loss','Adam and binary cross-entropy'],['Class weights','Crop 2.847, Weed 0.607'],['Random seed','42'],['Hardware','AMD Ryzen 3 5300U, CPU execution']], [2.6,3.5]),
        H('Why class weights'),
        P('There are far fewer crop images. Balanced training class weights make a crop example contribute more to loss than a weed example. They are calculated only from training labels. They do not create more images, alter the held-out label counts, or change the ordinary accuracy calculation.'),
        H('Why only five epochs'),
        P('This is a stated local CPU run budget, not a claim that five is the optimum number. The code supports up to ten head epochs and five fine-tuning epochs. Larger budgets take longer and should be evaluated using validation performance.'))
    add('9 Loss Accuracy and Evaluation Metrics',
        H('Loss and accuracy are different'),
        P('Accuracy is the fraction of predictions with the correct class after applying the 0.5 threshold. Binary cross-entropy loss also considers the score’s strength. A confidently wrong score receives a larger penalty than a less confident wrong score. Two models can therefore have similar accuracy but different loss.'),
        H('Why accuracy alone is insufficient'),
        P('The test set contains 685 weeds and 146 crops. A rule that always says Weed gets 685 of 831 images correct: 82.43% accuracy. Yet its Crop recall is zero. This is a calculated baseline, not your VGG16 result. Class-specific metrics reveal whether the model actually recognises crops.'),
        H('Precision and recall'),
        P('Crop precision asks: of the images predicted as Crop, how many really are crops? Crop recall asks: of all actual crop images, how many were identified as Crop? The same definitions apply separately to Weed. A crop predicted as Weed lowers Crop recall and could matter in a future weed-removal system.'),
        H('F1 score and support'),
        P('F1 combines precision and recall using their harmonic mean: twice their product divided by their sum. It is high only when both are reasonably high. Support is the number of actual images for a class: 146 for Crop and 685 for Weed in this test set.'),
        P('A macro average gives each class equal weight. A weighted average uses class support, so the larger Weed class contributes more. This evaluation average is different from the balanced class weights used in training loss.'),
        H('Confusion matrix'),
        P('The matrix counts the four possible outcomes. Rows are actual labels and columns are predicted labels. The diagonal cells are correct predictions. The off-diagonal cells are mistakes. Read the row and column labels before interpreting any matrix.'),
        H('ROC AUC'),
        P('ROC means Receiver Operating Characteristic, and AUC means area under the curve. ROC AUC summarises how well scores rank Weed images above Crop images across thresholds. AUC of 0.5 corresponds to chance-like ranking, while 1.0 is perfect ranking on the evaluated data. It differs from accuracy at the fixed 0.5 threshold.'))
    epoch_rows=[['Epoch','Phase','Train accuracy','Validation accuracy']]
    if history:
        for i,(train,val) in enumerate(zip(history['accuracy'],history['val_accuracy'])):
            epoch_rows.append([str(i+1),'Head' if i<(r['head_epochs_requested'] if r else 3) else 'Block5',pct(train),pct(val)])
    else:epoch_rows.append(['Pending','Internal draft','Pending','Pending'])
    add('10 Read the Accuracy Graph',
        P('Open presentation slide 7. The horizontal axis is epoch number. The vertical axis is accuracy. The blue line is training accuracy and the orange line is validation accuracy. Both lines are measured from this run; the test accuracy is a separate final number.'),
        I('accuracy_graph.png'),T(epoch_rows,[.65,.9,2.2,2.35]),
        H('What to look for'),
        P('Increasing training accuracy means the model is fitting training examples better. Validation accuracy describes performance on its separate validation images. A growing training-versus-validation gap can suggest overfitting, but one fluctuation alone is not enough to diagnose it.'),
        P('Validation accuracy can be higher than training accuracy because training uses augmentation and dropout, which make the training task harder. Validation uses unaugmented images and dropout is disabled. Epochs one through three train the head; epochs four and five fine-tune block5.'))
    best_val=min(history['val_loss']) if history else None
    add('11 Read the Loss Graph and Checkpoints',
        P('Open presentation slide 8. Lower binary cross-entropy is generally better for the same evaluation setting. The blue curve is weighted training loss and the orange curve is unweighted validation loss, so their magnitudes are not directly interchangeable.'),
        I('loss_graph.png'),
        H('What the callbacks do'),
        P('A callback is an action performed during training. ModelCheckpoint saves the model when validation loss improves. EarlyStopping can stop a phase after three epochs without improvement and restore its best weights. ReduceLROnPlateau halves the learning rate after two epochs without improvement. With a three-plus-two budget, early stopping may not trigger.'),
        P((f'The lowest measured validation loss across both phases was {best_val:.4f}. The test evaluation reloads the checkpoint selected using validation loss, rather than automatically using the last epoch.' if history else 'The final guide will insert the lowest measured validation loss here after training.')),
        H('How to explain an increase in validation loss'),
        P('Training is not guaranteed to improve every epoch. An increase can reflect less useful updates, overfitting or ordinary variation. Check the full curves and class metrics. In this run, the second head epoch had higher validation loss than the first, so the first checkpoint stayed eligible as the better model.'),
        P('Checkpoint selection never uses the test set. This separation is what makes the final test result a more honest estimate than repeatedly selecting models on test accuracy.'))
    if measured:
        m=r['confusion_matrix'];report=r['classification_report'];total=r['evaluated_test_images'];correct=m[0][0]+m[1][1]
        result_blocks=[
            P(summary),
            T([['Metric','Measured value'],['Correct predictions',f'{correct} of {total}'],['Test accuracy',pct(r['test_accuracy'])],['Test loss',f"{r['test_loss']:.4f}"],['ROC AUC',f"{r['auc']:.4f}"],['Always-Weed baseline','82.43% accuracy, 0% Crop recall']], [3.0,3.1]),
            H('Read the four counts'),
            T([['Actual / Predicted','Crop','Weed'],['Crop',str(m[0][0]),str(m[0][1])],['Weed',str(m[1][0]),str(m[1][1])]], [3.3,1.4,1.4]),
            P(f"Of 146 actual crop images, {m[0][0]} were correctly identified and {m[0][1]} were predicted as Weed. Of 685 actual weed images, {m[1][1]} were correctly identified and {m[1][0]} were predicted as Crop."),
            T([['Class','Precision','Recall','F1'],*[ [label,*[pct(report[label][k]) for k in ['precision','recall','f1-score']]] for label in ['Crop','Weed'] ]],[1.5,1.5,1.55,1.55]),
            P(f"Read overall accuracy together with the class metrics. Crop recall is {pct(report['Crop']['recall'])} and Weed recall is {pct(report['Weed']['recall'])}. In this split, about one in ten crop images was misclassified, even though total accuracy was high. That gap matters when considering future crop protection."),
            H('The appropriate conclusion'),
            P('These are measured full-dataset results on the saved image split. They are not smoke-test scores. The split excludes identical paths, but repeated plant identities and controlled backgrounds remain limitations. A good test score here does not establish reliable performance in a field.')]
    else:result_blocks=[P('Internal draft only. This page will contain the actual test metrics, confusion matrix and class-specific results after full training completes.')]
    add('12 Understand Your Actual Test Results',*result_blocks)
    add('13 Understand the Code Without Memorising It',
        H('Four Python ideas'),
        P('A variable names a value, such as batch_size. A function is a named reusable action, such as predict_image. A class groups related data and actions; Experiment stores the datasets, model and results. An import loads definitions from a library or another project file.'),
        C('experiment = Experiment(config)\nmodel = experiment.build()\nhead_history = experiment.train_head()'),
        P('The first line creates an Experiment object using the configuration. The second builds its model. The third starts head training and stores the returned history in a variable. The equals sign assigns a value; it is not a mathematical claim that both sides were already equal.'),
        T([['File or method','Purpose'],['scripts/dataset.py','Finds species images, maps labels and checks dataset counts.'],['Experiment.explore()','Verifies the dataset and saves distribution and sample plots.'],['Experiment.prepare()','Creates and saves disjoint splits, batches and class weights.'],['Experiment.build()','Loads ImageNet VGG16, adds the head and checks preprocessing.'],['train_head() and fine_tune()','Run the two phases with validation-based callbacks.'],['plot_history() and evaluate()','Save curves, reload the checkpoint, measure the test set and save predictions.'],['scripts/train_vgg16.py','Command-line entry point for the shared pipeline.'],['Main Jupyter notebook','Explains and runs the same pipeline in Jupyter.']], [2.5,3.6]),
        C('notebooks/VGG16_Crop_Weed_Classification.ipynb'),
        H('What was tested'),
        P('A small smoke run checked that the code executes, weights update, preprocessing is correct, and a saved model reloads consistently. A separate test checked that block5 updates while block1 remains frozen. All 25 notebook code cells executed in a smoke copy. Full-dataset training was run through the shared command-line pipeline; its measured results are the ones used in the final artifacts.'),
        P('A smoke test is a quick correctness check on a tiny subset. It is not evidence of full-dataset model accuracy. The main notebook remains configured for normal full training.'))
    image_path=r.get('_example_image','data/raw/plant_seedlings/Maize/example.png') if measured else 'data/raw/plant_seedlings/Maize/example.png'
    add('14 Open and Run the Project',
        H('Open PowerShell in the project folder'),
        C("cd 'C:\\Users\\shrav\\Desktop\\dl study proect\\CNN-Crop-and-Weed-Classification'",'powershell'),
        P('The environment and dataset are already installed on this laptop. Use the repository’s Python executable below so the correct libraries are loaded.'),
        H('Open the notebook'),
        C('.\\.venv\\Scripts\\python.exe -m jupyter lab','powershell'),
        P('Select the Crop Weed VGG16 kernel. Open notebooks/Review_Completed_Results.ipynb and Run All to inspect the saved results without training. The main VGG16_Crop_Weed_Classification.ipynb explains the complete training pipeline; its Run All starts another experiment with maximum budgets of ten head epochs and five fine-tuning epochs.'),
        H('Predict one image using the saved model'),
        C('.\\.venv\\Scripts\\python.exe scripts/predict_image.py `\n  --image "'+image_path+'"','powershell'),
        P('This example uses an actual saved test-image path. The command prints Crop or Weed and the Weed probability. Replace the quoted image path with your own file path to try another image. Prediction loads the saved model and does not retrain it.'),
        H('Repeat the same full training budget if needed'),
        C('.\\.venv\\Scripts\\python.exe scripts/train_vgg16.py `\n  --batch-size 4 --threads 4 `\n  --head-epochs 3 --fine-epochs 2','powershell'),
        P('The backtick continues a PowerShell command on the next line. Re-running training writes a new checkpoint and new output files. CPU training is slow, so you do not need to rerun it just to read results or present the project.'),
        H('Where the completed results are'),
        P('Open outputs/results_summary.json for the metric summary, outputs/training_history.json for epoch values, and outputs/test_predictions.csv for every test prediction. The .keras checkpoint is under models/. The revised presentation, report and this guide have Results or Explained in their names.'))
    add('15 What to Say for Slides 1 to 6',
        P('Use these short explanations as practice notes. Understand each idea before trying to memorise the wording. The presentation’s speaker notes allocate ten minutes in total.'),
        H('Slide 1 Project title'),
        P('My individual project is Crop and Weed Classification using VGG16. I used ImageNet transfer learning and evaluated the model on separate plant seedling test images.'),
        H('Slide 2 Problem statement'),
        P('The task is to take one RGB plant image and predict Crop or Weed. Young plants can look similar, and the dataset contains many more weeds than crops. The objective is to adapt visual features and measure performance fairly.'),
        H('Slide 3 Dataset'),
        P('The V2 Plant Seedlings dataset has 5,539 images and twelve species. Maize, Common Wheat and Sugar Beet form the Crop class, with 973 images. The other nine species form Weed, with 4,566 images. These photographs are original dataset samples.'),
        H('Slide 4 Preprocessing and split'),
        P('I split the data before augmentation using seed 42 and binary stratification. There are 3,877 training images, 831 validation images and 831 test images. Images are resized to 224 by 224. VGG16 preprocessing uses raw 0–255 pixels, channel reversal and mean subtraction.'),
        H('Slide 5 Model architecture'),
        P('The ImageNet VGG16 base extracts feature maps. Global average pooling reduces them to 512 features. A Dense 128 layer with ReLU and Dropout 0.5 learns the task. One sigmoid output gives the Weed score, with a threshold of 0.5.'),
        H('Slide 6 Hyperparameters'),
        P('The measured local run used batch size four. First I trained the head for three epochs with the base frozen at learning rate 0.001. Then I fine-tuned block5 for two epochs at 0.00001. Class weights address imbalance, and validation loss chooses the saved checkpoint.'))
    if measured:
        accuracy_statement=f"The measured test accuracy is {pct(r['test_accuracy'])} on 831 held-out images."
        loss_statement=f"The measured test loss is {r['test_loss']:.4f}, and ROC AUC is {r['auc']:.4f}."
        matrix_statement=f"{m[0][1]} crops were predicted as Weed, and {m[1][0]} weeds were predicted as Crop."
    else:accuracy_statement=loss_statement=matrix_statement='Internal draft: actual measurements will be inserted.'
    add('16 What to Say for Slides 7 to 11',
        H('Slide 7 Accuracy results'),
        P('The curves show training and validation accuracy across all five epochs. Blue is training and orange is validation. '+accuracy_statement+' The test set was not used to select the model.'),
        H('Slide 8 Loss results'),
        P('The loss graph shows binary cross-entropy. Lower validation loss determines the best checkpoint. '+loss_statement+' Training loss uses class weights, while validation loss is unweighted, so I interpret trends rather than assuming the two values should match.'),
        H('Slide 9 Confusion matrix and class metrics'),
        P('Rows are actual labels and columns are predicted labels. Diagonal cells are correct predictions. '+matrix_statement+' Precision, recall and F1 describe each class separately, which is important because Weed is the majority class.'),
        H('Slide 10 Sample predictions'),
        P('These are actual test photographs with the true label, predicted label and saved Weed score. I included correct examples and errors. The score explains the thresholded decision, but it is not a calibrated guarantee of correctness.'),
        H('Slide 11 Conclusion'),
        P('The project implements a complete local transfer-learning pipeline and reports measured test results. I also compare accuracy with the 82.43% majority-Weed baseline. Controlled backgrounds and repeated plant identities limit generalisation. Future work should verify plant-level splits and evaluate real field images.'),
        H('If you forget a technical term'),
        P('Return to the simple story: images become features, the head learns two labels, validation selects a model, and separate test images measure the result. Explain the idea in plain words before giving the term.'),
        H('Claims to state accurately'),
        P('This run used pretrained ImageNet weights and five measured epochs. It was executed on the CPU, and the split is at image level. Explain those facts directly rather than saying that every layer was trained from scratch or that field reliability has already been established.'))
    add('17 Viva Questions About the Model',
        H('1 What is the aim of your project'),
        P('To classify a seedling photograph as Crop or Weed using VGG16 transfer learning and evaluate the selected model on a separate test partition.'),
        H('2 Why did you use a CNN'),
        P('CNNs learn spatial patterns from images using shared filters. This task is image classification. The uploaded LSTM and GRU decks were layout references; those sequence models are not the architecture used in this project.'),
        H('3 Why VGG16 and ImageNet weights'),
        P('VGG16 provides an established pretrained visual feature extractor. ImageNet weights give a useful starting point when the plant dataset is much smaller. I then train a head for the new binary task.'),
        H('4 Why remove the original top layers'),
        P('The original classifier predicts ImageNet categories. include_top=False removes it so I can attach global average pooling and a new Crop/Weed head.'),
        H('5 Why one sigmoid output'),
        P('There are two mutually exclusive classes. One sigmoid score represents Weed, with Crop selected below the threshold. Binary cross-entropy is appropriate for this label encoding.'),
        H('6 Why freeze and then fine tune'),
        P('Freezing lets the new head learn without immediately changing pretrained features. Fine-tuning block5 later adapts high-level features using a smaller learning rate, while earlier blocks stay fixed.'),
        H('7 Why Dropout 0.5 and global average pooling'),
        P('Dropout reduces dependence on particular head activations during training. Global average pooling creates a compact 512-feature vector and uses fewer new parameters than flattening all feature-map positions.'))
    add('18 Viva Questions About Data and Results',
        H('8 Why have validation and test sets'),
        P('Validation guides checkpoint selection and training decisions. Test images are held out until the selected model is evaluated. Using test performance to choose a model would make the final estimate less fair.'),
        H('9 Why not divide image values by 255'),
        P('This VGG16 preprocessing convention expects raw 0–255 pixels before reversing channels and subtracting ImageNet means. Dividing first changes the intended input scale.'),
        H('10 What are class weights and why use them'),
        P('They change training loss contributions so minority crop examples matter more. The weights are calculated from training labels only. They do not change the validation or test distributions.'),
        H('11 Is high accuracy enough'),
        P('No. Always predicting Weed would already get 82.43% accuracy on this test split while missing every crop. I therefore inspect Crop recall, Weed recall, precision, F1 and the confusion matrix.'),
        H('12 What does five epochs mean'),
        P('The model passes through the full training partition five times: three frozen-head epochs and two block5 epochs. It is the stated CPU budget for this measured run. The default code permits larger maximum budgets.'),
        H('13 Are your results smoke test results'),
        P('No. Smoke tests checked execution on tiny balanced subsets. The final metrics use 831 test images from the full 5,539-image dataset, after full-partition training.'),
        H('14 What is the main limitation and future work'),
        P('Different image paths may still show the same plant at different growth stages, and the backgrounds are controlled. Future work should verify plant identities for grouped splits, obtain field photographs and evaluate threshold choices on validation data.'))
    add('19 Glossary and Reliable Reading',
        T([['Term','Plain meaning'],['Dataset','The collection of labelled examples.'],['Feature','A numerical representation of an image pattern.'],['Parameter or weight','A model value learned during training.'],['Hyperparameter','A setting chosen before training, such as batch size.'],['Epoch','One pass through the training partition.'],['Batch','A small group processed in one training step.'],['Optimizer','The method that updates trainable weights.'],['Checkpoint','A saved model selected during training.'],['Inference','Using learned weights to predict an image.'],['Overfitting','Fitting training-specific patterns that fail to generalise well.'],['Kernel in Jupyter','The Python process used to execute notebook code.'],['.venv','The isolated Python environment for this repository.'],['.ipynb','A notebook file containing explanation and code cells.'],['.keras','The saved model file used for later predictions.'],['CSV and JSON','Readable formats that store tables and structured results.']], [2.0,4.1]),
        H('Sources to read after this guide'),
        L('TensorFlow tutorial on transfer learning and fine tuning','https://www.tensorflow.org/tutorials/images/transfer_learning'),
        L('TensorFlow VGG16 preprocessing documentation','https://www.tensorflow.org/api_docs/python/tf/keras/applications/vgg16/preprocess_input'),
        L('Simonyan and Zisserman VGG paper','https://arxiv.org/abs/1409.1556'),
        L('Giselsson and colleagues Plant Seedlings dataset paper','https://arxiv.org/abs/1711.05458'),
        P('Project-specific facts come from scripts/core.py, outputs/dataset_splits.csv, outputs/training_history.json, outputs/results_summary.json and outputs/test_predictions.csv. The presentation uses measured output files, not the numbers in the two reference presentations.'))
    return pages

def format_table(doc,rows,widths):
    table=doc.add_table(rows=0,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for i,values in enumerate(rows):
        cells=table.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            if widths:cell.width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(value)
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17324D' if i==0 else ('FFFFFF' if i%2 else 'F1F4F7'));props.append(shade)
            margins=OxmlElement('w:tcMar')
            for name,val in [('top',65),('bottom',65),('left',100),('right',100)]:
                x=OxmlElement('w:'+name);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');margins.append(x)
            props.append(margins)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
            props.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
        no_split=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(no_split)
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    if widths:
        for col,w in zip(table.columns,widths):col.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def build(pages,destination,markdown):
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',27),('Heading 1',21),('Heading 2',13)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(8)
        for border in list(st._element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
        for fonts in st._element.findall('.//'+qn('w:rFonts')):
            for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+attribute),None)
    head=s.header.paragraphs[0];head.text='VGG16 Crop and Weed Project Study Guide'
    head.runs[0].font.size=Pt(9);head.runs[0].font.color.rgb=RGBColor(0,0,0)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run('Nalacherla Shravan Kumar   23EG110C43   Page ').font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    md=[]
    for index,(title,blocks) in enumerate(pages):
        title_paragraph=doc.add_paragraph(title,'Title' if index==0 else 'Heading 1')
        if index:title_paragraph.paragraph_format.page_break_before=True
        md.append('# '+title+'\n')
        for block in blocks:
            kind=block[0]
            if kind=='p':doc.add_paragraph(block[1]);md.append(block[1]+'\n')
            elif kind=='h':doc.add_paragraph(block[1],'Heading 2');md.append('## '+block[1]+'\n')
            elif kind=='code':
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
                run=p.add_run(block[1]);run.font.name='Consolas';run.font.size=Pt(9)
                md.append('```'+block[2]+'\n'+block[1]+'\n```\n')
            elif kind=='table':
                rows,widths=block[1:];format_table(doc,rows,widths)
                md.append('| '+' | '.join(rows[0])+' |\n| '+' | '.join(['---']*len(rows[0]))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows[1:])+'\n')
            elif kind=='image':
                path=ROOT/'outputs'/block[1]
                if path.exists():
                    doc.add_picture(str(path),width=Inches(6.15))
                    md.append(f'![Measured {block[1].replace("_graph.png","")} graph](../outputs/{block[1]})\n')
            elif kind=='link':
                p=doc.add_paragraph();link=OxmlElement('w:hyperlink')
                rel=doc.part.relate_to(block[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
                link.set(qn('r:id'),rel);run=OxmlElement('w:r');props=OxmlElement('w:rPr')
                color=OxmlElement('w:color');color.set(qn('w:val'),'17324D');props.append(color)
                run.append(props);text=OxmlElement('w:t');text.text=block[1];run.append(text);link.append(run);p._p.append(link)
                md.append(f'[{block[1]}]({block[2]})\n')
    settings=doc.settings._element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    destination.parent.mkdir(exist_ok=True,parents=True);doc.save(destination)
    markdown.parent.mkdir(exist_ok=True,parents=True);markdown.write_text('\n'.join(md),encoding='utf8')
    print(f'Guide built with {len(pages)} planned pages: {destination}')

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--draft',action='store_true');args=parser.parse_args()
    source=ROOT/'outputs/results_summary.json'
    r=json.loads(source.read_text()) if source.exists() else None
    if r and r.get('status')!='full_training_completed':r=None
    if not args.draft and (not r or r.get('status')!='full_training_completed'):raise RuntimeError('Completed full-run results are required.')
    history_path=ROOT/'outputs/training_history.json'
    h=json.loads(history_path.read_text()) if history_path.exists() else None
    if r:
        import csv
        with (ROOT/'outputs/test_predictions.csv').open(encoding='utf8',newline='') as f:
            predictions=list(csv.DictReader(f))
        r['_example_image']=next(row['relative_path'] for row in predictions if row['label']=='0')
    split=None
    pages=pages_for(r,h,split)
    destination=ROOT/('.build/guide-draft.docx' if args.draft else 'docs/Roll_43_Project_Explained_From_Zero.docx')
    markdown=ROOT/('.build/guide-draft.md' if args.draft else 'docs/Project_Explained_From_Zero.md')
    build(pages,destination,markdown)

if __name__=='__main__':main()
