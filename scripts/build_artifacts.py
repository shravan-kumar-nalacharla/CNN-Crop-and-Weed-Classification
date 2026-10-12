"""Rebuild the shared-core notebook and reference-based individual report."""
import json
import io
import zipfile
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
TITLE='CNN Crop and Weed Classification using VGG16'

def notebook():
    class NotebookJSON:
        @staticmethod
        def new_markdown_cell(source):
            return {'cell_type':'markdown','metadata':{},'source':source.splitlines(keepends=True)}
        @staticmethod
        def new_code_cell(source):
            return {'cell_type':'code','metadata':{},'source':source.splitlines(keepends=True),'execution_count':None,'outputs':[]}
        @staticmethod
        def new_notebook(cells,metadata):
            import uuid
            for cell in cells:cell['id']=uuid.uuid4().hex[:8]
            return {'nbformat':4,'nbformat_minor':5,'cells':cells,'metadata':metadata}
        @staticmethod
        def write(nb,path):path.write_text(json.dumps(nb,indent=2),encoding='utf-8')
    class API:
        v4=NotebookJSON
        write=NotebookJSON.write
    n=API
    cells=[]
    def add(title, prose, code=None):
        cells.append(n.v4.new_markdown_cell(f'## {title}\n\n{prose}'))
        if code: cells.append(n.v4.new_code_cell(code))
    cells.append(n.v4.new_markdown_cell(f'# {TITLE}\n\n**Nalacherla Shravan Kumar • 23EG110C43 • Roll Number 43**\n\nIndividual study project. ImageNet transfer learning, local Jupyter, 12 species mapped to Crop (0) and Weed (1). Full training is the default. Smoke metrics are execution checks only.'))
    bootstrap="""from pathlib import Path
import sys, os
ROOT = next((p for p in [Path.cwd(), *Path.cwd().parents] if (p/'scripts/core.py').exists()), None)
assert ROOT is not None, 'Start Jupyter from this repository or its notebooks folder.'
sys.path.insert(0, str(ROOT/'scripts'))
from verify_environment import verify
environment = verify()
"""
    add('Environment verification','Select the Crop Weed VGG16 kernel. Dependencies must be installed in .venv.',bootstrap)
    add('Configuration','Batch 16 by default. Set SMOKE_TEST=True for one quick epoch, or reduce batch size to 2 for low-memory laptops.',"from core import Config, Experiment, predict_image\nSMOKE_TEST = False\nconfig = Config(batch_size=16, smoke_test=SMOKE_TEST)\nexperiment = Experiment(config)")
    add('Dataset location','Download first with python scripts/download_dataset.py. Nested extraction is supported.',"print(config.data_dir)")
    add('Species validation','Decode every image and confirm the 12 expected species and 5,539 exact V2 counts.',"from dataset import SPECIES, CROPS\ndf = experiment.explore()\ndisplay(df.groupby('species').size().to_frame('images'))")
    add('Binary label mapping','Crop: Maize, Common Wheat, Sugar Beet. The other nine species are Weed.',"display(df[['species','label']].drop_duplicates().sort_values('label'))\nassert set(df.label) == {0, 1}")
    add('Dataset exploration','Inspect paths and image counts. Images use controlled backgrounds and may share plant identity across growth stages.',"display(df.head())\nprint(df.shape)")
    show=lambda name: f"from IPython.display import Image, display\ndisplay(Image(filename=str(experiment.output/'{name}.png')))"
    add('Class distribution','Class imbalance motivates training-only balanced weights.',show('class_distribution'))
    add('Sample images','Inspect representative original RGB images.',show('sample_images'))
    add('Stratified splitting','Seed 42. Split before augmentation, 70% training, 15% validation, 15% test. Assert no overlapping paths. Save dataset_splits.csv.',"splits = experiment.prepare()\ndisplay(splits.groupby(['split','label']).size().unstack())")
    add('TensorFlow data pipeline','Load and resize raw RGB pixels to 224×224. Keep values in 0–255. A small prefetch queue limits memory.',"images, labels = next(iter(experiment.train_ds))\nprint(images.shape, labels.shape, float(images.numpy().max()))")
    add('Data augmentation','Training-only horizontal flip, small rotation, zoom and translation. Keras disables random transforms during evaluation.',"from core import augmentation\naugmentation(config.seed).summary()")
    add('Class weights','Balanced weights are calculated only from training labels.',"print(experiment.weights)")
    add('VGG16 architecture','ImageNet VGG16 without top; preprocessing after augmentation; global average pooling; Dense 128 ReLU; Dropout 0.5; one sigmoid.',"model = experiment.build()")
    add('Phase 1 training','Freeze all VGG16 layers. Adam 0.001, binary cross-entropy, at most 10 epochs. Checkpoint, early stopping and plateau reduction monitor validation loss.',"head_history = experiment.train_head()")
    add('Phase 2 fine-tuning','Unfreeze only block5, recompile Adam at 0.00001 and train up to five additional epochs. Smoke mode validates flags without a second epoch.',"fine_history = experiment.fine_tune()")
    add('Combined training history','Combine both phases and save the measured epoch history.',"history = experiment.plot_history()\ndisplay(history)")
    add('Accuracy graph','Compare training and validation accuracy by epoch.',show('accuracy_graph'))
    add('Loss graph','Validation loss selects the best checkpoint across both phases.',show('loss_graph'))
    add('Test evaluation','Reload the best checkpoint. Verify save/reload prediction equivalence. Evaluate the held-out test images once.',"results = experiment.evaluate()\nprint('Test loss:', results['test_loss'])\nprint('Test accuracy:', results['test_accuracy'])\nprint('AUC:', results['auc'])")
    add('Confusion matrix','Rows are actual labels; columns are predicted labels. Crop=0, Weed=1.',show('confusion_matrix'))
    add('Classification report','Both classes include precision, recall, F1 and support. Undefined precision is reported as zero.',"import pandas as pd\ndisplay(pd.read_csv(experiment.output/'classification_report.csv', index_col=0))")
    add('Sample predictions','The sigmoid gives Weed probability, thresholded at 0.5.',show('sample_predictions'))
    add('Misclassified examples','Inspect errors before drawing conclusions about field use.',show('misclassified_examples'))
    add('Single-image prediction','Pass an RGB image path. Saved model includes VGG16 preprocessing.',"example_path = experiment.test.iloc[0].path\npredict_image(experiment.model, example_path)")
    add('Saving model and results','All outputs use repository-relative directories. Smoke results remain separate. Rebuild report and presentation after a full run.',"print('Best model:', experiment.model_path.relative_to(ROOT))\nprint('Results:', (experiment.output/'results_summary.json').relative_to(ROOT))\nassert experiment.model_path.exists()\nassert (experiment.output/'test_predictions.csv').exists()")
    nb=n.v4.new_notebook(cells=cells,metadata={'kernelspec':{'name':'crop-weed-vgg16','display_name':'Crop Weed VGG16','language':'python'},'language_info':{'name':'python','version':'3.11'}})
    (ROOT/'notebooks').mkdir(exist_ok=True)
    n.write(nb,ROOT/'notebooks/VGG16_Crop_Weed_Classification.ipynb')

def result_status():
    p=ROOT/'outputs/results_summary.json'
    if p.exists():
        data=json.loads(p.read_text())
        if data.get('status')=='full_training_completed': return data
    return None

def report():
    reference=ROOT/'references/Sample report.docx'
    doc=Document(reference)
    logo = next((deepcopy(p._p) for p in doc.paragraphs if p._p.xpath('.//w:drawing')), None)
    # Retain reference section geometry, styles, relationships and page furniture.
    body=doc._element.body
    for element in list(body):
        if element.tag != qn('w:sectPr'): body.remove(element)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        doc.styles[name].font.name='Times New Roman'
        fonts=doc.styles[name]._element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+attr),None)
            fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
    result=result_status()
    pending='PENDING FULL TRAINING: insert measured results from outputs/results_summary.json after the complete run. Smoke-test metrics are not final performance.'
    def p(text):
        paragraph=doc.add_paragraph(text);paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        return paragraph
    def heading(text): doc.add_heading(text,level=1)
    title=doc.add_paragraph(TITLE,'Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for text in ['Deep Learning Study Project Report','Individual Project • Roll Number 43','B.Tech IV Year I Semester • CSE DS','Nalacherla Shravan Kumar','23EG110C43','Under Guidance of Dr P Rajeshwari','Department of Computer Science Engineering Data Science','ANURAG UNIVERSITY','Venkatapur, Ghatkesar, Medchal Malkajgiri, Telangana 500088','One-sided printing • October 2026']:
        if text.startswith('Department') and logo is not None:
            body.insert(len(body)-1,logo)
        paragraph=doc.add_paragraph(text);paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    pages=[
      [('Abstract','This individual project develops a binary Crop and Weed classifier using ImageNet-pretrained VGG16. The V2 Plant Seedlings dataset supplies twelve species mapped to two classes. A stratified image-level split separates training, validation and testing before augmentation. The method first trains a new classification head, then fine-tunes only the final convolution block using a smaller learning rate. Balanced class weights address the majority Weed class. A shared Python core drives a local Jupyter notebook and command-line runner. Held-out evaluation reports class-specific precision, recall, F1, AUC and a confusion matrix. '+('Measured full-run results appear in the performance section.' if result else 'Full-run results remain pending; a smoke test establishes execution correctness only.')),
       ('Introduction','Crop and weed discrimination can support selective weed management. Visually similar seedlings make this a challenging image-classification task. Transfer learning adapts previously learned image features to a smaller agricultural dataset. This project aligns with the syllabus units on optimization, convolutional networks and transfer learning.')],
      [('Problem statement','Given one RGB image, predict Crop or Weed according to the defined species mapping. Controlled dataset accuracy does not establish reliable field deployment.'),('Objectives','Implement an individual VGG16 project for Roll 43; reproduce data validation and nonoverlapping splits; compare training and validation behavior; evaluate the held-out test set; provide local setup, notebook, model and documented outputs.')],
      [('Dataset description','V2 Plant Seedlings contains 5,539 RGB images from twelve species. The original source describes approximately 960 unique plants recorded at multiple growth stages. Source: Kaggle vbookshelf/v2-plant-seedlings-dataset and Aarhus University. License: CC BY-SA 4.0. Crop contains Maize, Common Wheat and Sugar Beet (973 images). Weed contains the remaining nine species (4,566 images). The verifier decodes all files and checks every species count. Folder nesting and apostrophe variants are normalized.'),('Data preprocessing','Recursively identify species from parent folders, validate image readability and convert to RGB. Resize to 224×224. Keep float pixel values in 0–255; VGG16 preprocessing converts RGB to BGR and subtracts ImageNet means. Split before augmentation, using seed 42 and binary stratification. Training-only transforms apply horizontal flips, rotation 0.05, zoom 0.1 and translation 0.05.')],
      [('VGG16 architecture','The ImageNet convolution base contains 13 convolutional layers grouped into five blocks, each followed by pooling. include_top=False removes the original classifier. The model contains input and augmentation, VGG16 preprocessing, VGG16 base, GlobalAveragePooling2D, Dense 128 ReLU, Dropout 0.5 and Dense 1 sigmoid. Sigmoid output estimates Weed probability and a 0.5 threshold selects the binary class. Assertions confirm the VGG16 structure, requested ImageNet weights, trainability and sigmoid output.'),('Framework and tools','Python 3.11, TensorFlow/Keras, NumPy, pandas, scikit-learn, Matplotlib, Seaborn and Pillow implement the experiment. Kaggle and kagglehub acquire data. JupyterLab and ipykernel support local notebook execution. A repository-specific .venv isolates dependencies. Native Windows uses CPU; WSL2 or optional Colab supports NVIDIA acceleration.')],
      [('Methodology','Verify environment and dataset, explore distributions, create stratified 70/15/15 partitions and assert path disjointness. Calculate balanced class weights only from training labels. Create a tf.data pipeline that loads raw RGB images without normalizing by 255. Train the classification head, fine-tune block5, reload the global best validation-loss checkpoint and evaluate test predictions once.'),('Hyperparameters','Input 224×224×3; batch size 16 configurable; seed 42; Dense 128 ReLU; Dropout 0.5; binary cross-entropy; Adam 0.001 for up to 10 head epochs; Adam 0.00001 for up to 5 block5 epochs. EarlyStopping patience 3 restores best weights; ReduceLROnPlateau patience 2 halves learning rate; checkpoint monitors validation loss.')],
      [('Training procedure','Phase 1 freezes all VGG16 layers and learns the binary head. Phase 2 enables only block5 layers and recompiles the model. The checkpoint retains the best validation loss across both phases. No augmentation runs on validation or test images. The quick smoke test uses balanced subsets drawn after splitting and runs one head epoch; it checks forward/backward computation, block5 flags, preprocessing, model reload and prediction generation.')],
      [('Model performance',pending if not result else f"Measured test loss: {result['test_loss']:.5f}. Test accuracy: {result['test_accuracy']:.5f}. AUC: {result['auc']:.5f}. Test images: {result['evaluated_test_images']}."),('Accuracy graph','Training and validation accuracy by epoch, with the fine-tuning boundary marked.'),('Loss graph','Training and validation binary cross-entropy; the best validation loss selects the model.')],
      [('Confusion matrix','Rows represent actual class and columns predicted class. Crop false negatives represent crop images predicted as weeds, a potentially harmful error in future agricultural use.'),('Analysis and results',pending if not result else 'Read the measured class report beside the confusion matrix. Compare minority Crop recall with Weed recall and compare train/validation curves for overfitting. The AUC measures ranking across thresholds; the reported class decisions use threshold 0.5.')],
      [('Sample predictions','Review correct and incorrect test predictions using the generated image grids. A confident sigmoid output is not a calibrated guarantee.'),('Conclusion','The project implements reproducible VGG16 transfer learning and local execution for an individual study. '+('Full-run measurements quantify performance on this held-out image split.' if result else 'Full-run performance remains pending and no accuracy claim is made.')),('Limitations and future work','Controlled lighting and backgrounds limit field generalization. Multiple images may depict the same plant at different growth stages, so disjoint image paths do not ensure plant-level independence. Future work should verify plant identities for grouped splitting, obtain field images, calibrate probabilities and examine Crop recall under varied thresholds. CPU fine-tuning can be slow and memory intensive.')],
      [('References','1. Supplied DL Syllabus, ML Deep Learning Lab Programs, List of study projects and Sample report (references/).\n2. Giselsson et al. (2017), A Public Image Database for Benchmark of Plant Seedling Classification Algorithms. https://arxiv.org/abs/1711.05458\n3. Aarhus Plant Seedlings Dataset. https://vision.eng.au.dk/plant-seedlings-dataset/\n4. Kaggle V2 dataset. https://www.kaggle.com/datasets/vbookshelf/v2-plant-seedlings-dataset\n5. Simonyan and Zisserman (2014), Very Deep Convolutional Networks for Large-Scale Image Recognition. https://arxiv.org/abs/1409.1556\n6. TensorFlow VGG16 and installation documentation. https://www.tensorflow.org/api_docs/python/tf/keras/applications/VGG16 and https://www.tensorflow.org/install/pip\n7. CC BY-SA 4.0. https://creativecommons.org/licenses/by-sa/4.0/')]
    ]
    heading('Index')
    table=doc.add_table(rows=1,cols=2);table.style='Table Grid'
    table.rows[0].cells[0].text='Topic';table.rows[0].cells[1].text='Page'
    for page_number,sections in enumerate(pages,3):
        for name,_ in sections:
            cells=table.add_row().cells;cells[0].text=name;cells[1].text=str(page_number)
            for cell in cells:
                for run in cell.paragraphs[0].runs:run.font.size=Pt(10)
    doc.add_page_break()
    graphs={'Accuracy graph':'accuracy_graph','Loss graph':'loss_graph','Confusion matrix':'confusion_matrix','Sample predictions':'sample_predictions'}
    for i,sections in enumerate(pages):
        for name,text in sections:
            heading(name)
            paragraph=p(text)
            if name=='References':paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
            if name=='Model performance' and result:
                t=doc.add_table(rows=1,cols=4);t.style='Table Grid'
                for c,text in zip(t.rows[0].cells,['Class','Precision','Recall','F1']):c.text=text
                for label in ['Crop','Weed']:
                    vals=result['classification_report'][label]
                    for c,text in zip(t.add_row().cells,[label,*[f'{vals[k]:.4f}' for k in ['precision','recall','f1-score']]]):c.text=text
            if name in graphs:
                image=ROOT/f'outputs/{graphs[name]}.png'
                if result and image.exists():doc.add_picture(str(image),width=Inches(4.4))
                else:p('[Figure pending full training]')
        if i<len(pages)-1:doc.add_page_break()
    # Page field supports opening and printing in Word.
    foot=doc.sections[0].footer.paragraphs[0]
    foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    settings=doc.settings._element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    (ROOT/'report').mkdir(exist_ok=True)
    destination=ROOT/'report/Roll_43_VGG16_Report.docx'
    doc.save(destination)
    # python-docx may serialize opaque relationship parts differently. Preserve them exactly.
    buffer=io.BytesIO()
    with zipfile.ZipFile(reference) as source, zipfile.ZipFile(destination) as generated, zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as output:
        for item in generated.infolist():
            keep=item.filename.startswith(('customXml/','word/theme/','word/media/'))
            output.writestr(item,source.read(item.filename) if keep and item.filename in source.namelist() else generated.read(item.filename))
    destination.write_bytes(buffer.getvalue())

if __name__=='__main__':
    notebook();report()
